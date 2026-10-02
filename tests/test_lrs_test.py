"""LRS WELL TEST SUMMARY sheet -> manual test (POST /api/lrs/parse).

The workbook built here follows the layout of the MPE-48 sheet of 2026-10-02:
a label cell, the value a few cells to the right, then a unit.
"""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook

from server.services.lrs_test import parse_lrs_sheet


def _sheet(overrides=None, drop=()):
    """Workbook bytes. ``overrides`` replaces a label's value; ``drop`` omits rows."""
    rows = [
        ("WELL#:", "MPE-48", None),
        ("TEST DATE:", datetime(2026, 10, 2), None),
        ("DURATION:", 12.0, "HOURS"),
        ("RESERVOIR:", "MPU Schrader Bluff", None),
        ("TEST LOCATION CODE:", "LRS Unit 6", None),
        ("CHOKE:", 0, "Bypass"),
        ("WELL HEAD PRESSURE:", 363, "PSI"),
        ("WELL HEAD TEMPERATURE:", 108, "DEG F"),
        ("SEPARATOR PRESSURE:", 260, "PSI"),
        ("AVERAGE SPIN OUT W/C", 84.5, "%"),
        ("POWER FLUID RATE:", 4732.0, "BLPD"),
        ("POWER FLUID PRESSURE:", 3353.8, "PSI"),
        ("LIFT GAS RATE:", 0.0, "MMSCFPD"),
        ("CORRECTED FORMATION OIL RATE", 1039.2, "STBOPD"),
        ("FORMATION WATER RATE", 745.4, "BWPD"),
        ("CORRECTED WATER CUT", 41.8, "%"),
        ("CORRECTED FORMATION FLUID RATE", 1784.6, "STBFPD"),
        ("CORRECTED FORMATION GAS RATE", 0.0, "MMSCFPD"),
        ("CORRECTED FORMATION GOR", 1.8, "SCF/STB"),
        ("CORRECTED FORMATION TGOR", 1.8, "SCF/STB"),
    ]
    book = Workbook()
    ws = book.active
    ws.cell(row=1, column=4, value="WELL TEST SUMMARY")
    r = 2
    for label, value, unit in rows:
        if label in drop:
            continue
        value = (overrides or {}).get(label, value)
        # header block sits in column A; the measurements start in column C
        col = 1 if r < 8 else 3
        ws.cell(row=r, column=col, value=label)
        ws.cell(row=r, column=col + 5, value=value)
        if unit:
            ws.cell(row=r, column=col + 6, value=unit)
        r += 1
    # two labels on one row, like the sheet's start date / start time line
    ws.cell(row=r, column=1, value="TEST START DATE:")
    ws.cell(row=r, column=3, value=datetime(2026, 10, 1))
    ws.cell(row=r, column=5, value="TEST START TIME")
    ws.cell(row=r, column=7, value="17:00")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_reads_the_measurements_by_label():
    got = parse_lrs_sheet(_sheet(), "E-48 LRS.xlsx")
    assert got["well"] == "MPE-48"
    assert got["test_date"] == "2026-10-02"
    assert got["location"] == "LRS Unit 6"
    assert got["hours"] == 12.0
    assert (got["oil"], got["water"], got["total_fluid"]) == (1039.2, 745.4, 1784.6)
    assert got["form_wc"] == pytest.approx(0.418)
    assert (got["pf_rate"], got["pf_press"], got["whp"]) == (4732.0, 3353.8, 363.0)
    assert got["gor"] == 1.8
    assert got["missing"] == []


def test_formation_water_is_not_confused_with_the_spin_out_cut():
    # "AVERAGE SPIN OUT W/C" (84.5 %, total cut with PF) must never be the WC.
    got = parse_lrs_sheet(_sheet(drop=("CORRECTED WATER CUT",)), "t.xlsx")
    assert got["form_wc"] == pytest.approx(745.4 / 1784.6)


def test_missing_rate_is_derived_and_missing_extras_are_reported():
    got = parse_lrs_sheet(
        _sheet(drop=("CORRECTED FORMATION FLUID RATE", "POWER FLUID RATE:", "TEST DATE:")), "t.xlsx"
    )
    assert got["total_fluid"] == pytest.approx(1784.6)
    assert got["pf_rate"] is None
    assert got["test_date"] == "2026-10-01"  # falls back to the start date
    assert "Power Fluid Rate" in got["missing"]


def test_text_numbers_and_field_style_well_names_are_accepted():
    got = parse_lrs_sheet(
        _sheet({"WELL#:": "MPU E-048", "POWER FLUID RATE:": "4,732", "TEST DATE:": "10/2/2026"}), "t.xlsx"
    )
    assert got["well"] == "MPE-48"
    assert got["pf_rate"] == 4732.0
    assert got["test_date"] == "2026-10-02"


@pytest.mark.parametrize(
    "blob, name, text",
    [
        (b"not a workbook", "t.xlsx", "not a readable Excel workbook"),
        (b"", "old.xls", "save the sheet as .xlsx"),
    ],
)
def test_unreadable_files_say_why(blob, name, text):
    with pytest.raises(ValueError, match=text):
        parse_lrs_sheet(blob, name)


def test_a_workbook_that_is_not_an_lrs_sheet_is_refused():
    book = Workbook()
    book.active["A1"] = "Foreman schedule"
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ValueError, match="no formation oil"):
        parse_lrs_sheet(out.getvalue(), "schedule.xlsx")


def test_endpoint_round_trip_and_422():
    from fastapi.testclient import TestClient

    from server.main import app

    client = TestClient(app)
    kind = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ok = client.post("/api/lrs/parse", files={"file": ("E-48.xlsx", _sheet(), kind)})
    assert ok.status_code == 200
    assert ok.json()["oil"] == 1039.2 and ok.json()["filename"] == "E-48.xlsx"
    bad = client.post("/api/lrs/parse", files={"file": ("junk.xlsx", b"junk", kind)})
    assert bad.status_code == 422
    assert "junk.xlsx" in bad.json()["detail"]["message"]
